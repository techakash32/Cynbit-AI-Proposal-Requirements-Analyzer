from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from threading import Lock
from typing import Any


OCR_RENDER_DPI = 200
OCR_NATIVE_TEXT_MIN_CHARACTERS = 40
MAX_OCR_PAGES = 30
MAX_OCR_RENDER_PIXELS = 20_000_000
_RAPIDOCR_INFERENCE_LOCK = Lock()


class DocumentExtractionError(ValueError):
    """A user-actionable document extraction limit or capability error."""


@dataclass
class ExtractedDocument:
    text: str
    page_blocks: list[dict[str, Any]]
    ocr_used: bool = False


@lru_cache(maxsize=1)
def _get_rapidocr_engine() -> Any:
    """Load the local RapidOCR CPU engine only when a scanned page needs it."""
    try:
        from rapidocr import RapidOCR
    except ImportError as exc:
        raise DocumentExtractionError("Scanned PDF OCR is not installed; install backend/requirements.txt and restart the API") from exc
    return RapidOCR()


def _ocr_page(page: Any, pymupdf: Any) -> list[str]:
    import numpy as np

    width_pixels = max(1, int(round(page.rect.width * OCR_RENDER_DPI / 72)))
    height_pixels = max(1, int(round(page.rect.height * OCR_RENDER_DPI / 72)))
    if width_pixels * height_pixels > MAX_OCR_RENDER_PIXELS:
        raise DocumentExtractionError("A scanned PDF page is too large for OCR; reduce its page size or resolution and try again")
    pixmap = page.get_pixmap(dpi=OCR_RENDER_DPI, colorspace=pymupdf.csRGB, alpha=False)
    image = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape((pixmap.height, pixmap.width, pixmap.n))
    with _RAPIDOCR_INFERENCE_LOCK:
        result = _get_rapidocr_engine()(image)
    recognized_lines = getattr(result, "txts", None) or ()
    return [line.strip() for line in recognized_lines if isinstance(line, str) and line.strip()]


def extract_document(content: bytes, filename: str, content_type: str) -> ExtractedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf" or content_type == "application/pdf":
        import pymupdf

        doc = pymupdf.open(stream=content, filetype="pdf")
        blocks: list[dict[str, Any]] = []
        pages: list[str] = []
        ocr_used = False
        ocr_page_count = 0
        try:
            for page_offset in range(doc.page_count):
                page_number = page_offset + 1
                page = doc.load_page(page_offset)
                native_result = page.get_text("text")
                native_text = native_result.strip() if isinstance(native_result, str) else ""
                native_blocks: list[str] = []
                for block in page.get_text("blocks"):
                    if not isinstance(block, (tuple, list)) or len(block) <= 4:
                        continue
                    block_text = block[4]
                    if isinstance(block_text, str) and block_text.strip():
                        native_blocks.append(block_text.strip())
                if native_text and not native_blocks:
                    native_blocks = [native_text]

                page_text_parts = [native_text] if native_text else []
                page_blocks = list(native_blocks)
                compact_native_text = "".join(native_text.split())
                if len(compact_native_text) < OCR_NATIVE_TEXT_MIN_CHARACTERS:
                    if ocr_page_count >= MAX_OCR_PAGES:
                        raise DocumentExtractionError(
                            f"This PDF needs OCR on more than {MAX_OCR_PAGES} pages; split it into smaller files and try again"
                        )
                    ocr_page_count += 1
                    ocr_used = True
                    ocr_lines = _ocr_page(page, pymupdf)
                    new_ocr_lines = [line for line in ocr_lines if line not in native_text]
                    page_text_parts.extend(new_ocr_lines)
                    page_blocks.extend(new_ocr_lines)

                page_text = "\n".join(page_text_parts).strip()
                pages.append(page_text)
                blocks.extend({"page": page_number, "text": value} for value in page_blocks if value)
        finally:
            doc.close()
        return ExtractedDocument("\n".join(pages).strip(), blocks, ocr_used=ocr_used)

    if suffix == ".docx" or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        from io import BytesIO
        from docx import Document

        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        text = "\n".join(paragraphs).strip()
        return ExtractedDocument(text, [{"page": None, "text": line} for line in paragraphs], ocr_used=False)

    if suffix in {".txt", ".md"} or content_type.startswith("text/"):
        text = content.decode("utf-8", errors="replace").strip()
        return ExtractedDocument(text, [{"page": None, "text": line} for line in text.splitlines() if line.strip()], ocr_used=False)

    raise ValueError("Unsupported document type; provide PDF, DOCX, or TXT")


def ground_evidence(evidence: str | None, source_text: str, source_page: int | None = None, page_blocks: list[dict[str, Any]] | None = None) -> tuple[str | None, int | None]:
    """Return only verbatim source evidence; unsupported model claims become null."""
    if not evidence or evidence not in source_text:
        return None, None
    if page_blocks:
        for block in page_blocks:
            if evidence in str(block.get("text", "")):
                page = block.get("page")
                return evidence, page if isinstance(page, int) and page >= 1 else None
        return evidence, None
    return evidence, source_page
