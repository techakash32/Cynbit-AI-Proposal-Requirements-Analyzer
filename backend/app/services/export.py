from __future__ import annotations

from io import BytesIO
from typing import Any
from xml.sax.saxutils import escape

EVIDENCE_NOTE = "Evidence note: Proposal text does not carry sentence-level source citations in this prototype. Treat document-specific statements as unverified; no supporting evidence is established unless the matching requirement or risk shows a verbatim snippet and source page. Review every claim before sharing."


def markdown(sections: list[dict[str, Any]]) -> str:
    body = "\n\n".join(f"## {section['heading']}\n\n{section['content']}" for section in sections)
    return f"> **Evidence note:** {EVIDENCE_NOTE}\n\n{body}\n"


def docx_bytes(sections: list[dict[str, Any]]) -> bytes:
    from docx import Document
    document = Document()
    document.add_paragraph(EVIDENCE_NOTE)
    for section in sections:
        document.add_heading(section["heading"], level=2)
        document.add_paragraph(section["content"])
    output = BytesIO(); document.save(output); return output.getvalue()


def pdf_bytes(sections: list[dict[str, Any]]) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    output = BytesIO(); document = SimpleDocTemplate(output, pagesize=letter)
    styles = getSampleStyleSheet(); story = [Paragraph(escape(EVIDENCE_NOTE), styles["BodyText"]), Spacer(1, 12)]
    for section in sections:
        content = escape(section["content"]).replace("\n", "<br/>")
        story.extend([Paragraph(escape(section["heading"]), styles["Heading2"]), Paragraph(content, styles["BodyText"]), Spacer(1, 12)])
    document.build(story); return output.getvalue()
